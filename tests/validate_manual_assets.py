#!/usr/bin/env python3
"""Validate order-import manual assets: CSV template, JSON examples, handbook sections."""

from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
HANDBOOK = ROOT / "docs" / "旺店通旗舰版-订单导入操作手册.md"
TEMPLATE = ROOT / "templates" / "订单导入模板.csv"
TEMPLATE_XLSX = ROOT / "templates" / "订单导入模板.xlsx"
RAW_PUSH = ROOT / "examples" / "raw_trade_push_self2.json"
TRADE_IMPORT = ROOT / "examples" / "trade_import_upload.json"

REQUIRED_CSV_HEADERS = [
    "原始单号",
    "原始子单号",
    "店铺名称",
    "收件人",
    "手机",
    "省",
    "市",
    "区",
    "详细地址",
    "商家编码",
    "数量",
    "单价",
    "应收金额",
]

REQUIRED_HANDBOOK_HEADINGS = [
    "导入方式怎么选",
    "Excel / CSV 批量导入",
    "手工建单",
    "开放接口：原始单推送",
    "开放接口：已完成订单推送",
    "常见问题",
    "上线检查清单",
]


def fail(message: str) -> None:
    print(f"FAIL: {message}")
    raise SystemExit(1)


def check_handbook() -> None:
    text = HANDBOOK.read_text(encoding="utf-8")
    if not text.strip():
        fail("handbook is empty")
    for heading in REQUIRED_HANDBOOK_HEADINGS:
        if heading not in text:
            fail(f"handbook missing section: {heading}")
    if "sales.RawTrade.pushSelf2" not in text:
        fail("handbook missing pushSelf2 interface")
    if "sales.TradeImport.upload" not in text:
        fail("handbook missing TradeImport.upload interface")


def check_csv() -> None:
    raw = TEMPLATE.read_bytes()
    if not raw.startswith(b"\xef\xbb\xbf"):
        fail("CSV template must start with UTF-8 BOM for Excel")
    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(text.splitlines())
    headers = reader.fieldnames or []
    missing = [h for h in REQUIRED_CSV_HEADERS if h not in headers]
    if missing:
        fail(f"CSV missing headers: {missing}")

    rows = list(reader)
    if len(rows) < 2:
        fail("CSV should contain at least two sample rows")

    grouped: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        tid = (row.get("原始单号") or "").strip()
        if not tid:
            fail("CSV has empty 原始单号")
        if not (row.get("商家编码") or "").strip():
            fail(f"{tid}: empty 商家编码")
        qty = float(row["数量"])
        if qty <= 0:
            fail(f"{tid}: quantity must be > 0")
        grouped.setdefault(tid, []).append(row)

    for tid, items in grouped.items():
        first = items[0]
        for item in items[1:]:
            for key in ("收件人", "手机", "店铺名称", "省", "市", "区"):
                if item[key] != first[key]:
                    fail(f"{tid}: column {key} is inconsistent across lines")
        receivable_cells = [i["应收金额"].strip() for i in items if i["应收金额"].strip()]
        if len(receivable_cells) != 1:
            fail(f"{tid}: 应收金额 must appear on exactly one line")
        postage_cells = [i["邮费"].strip() for i in items if i["邮费"].strip()]
        if len(postage_cells) != 1:
            fail(f"{tid}: 邮费 must appear on exactly one line")

    try:
        from openpyxl import load_workbook
    except ImportError:
        if not TEMPLATE_XLSX.is_file():
            fail("missing Excel template")
        return
    if not TEMPLATE_XLSX.is_file():
        fail("missing Excel template")
    wb = load_workbook(TEMPLATE_XLSX)
    if "订单导入" not in wb.sheetnames or "填写说明" not in wb.sheetnames:
        fail("xlsx must contain 订单导入 and 填写说明 sheets")
    xheaders = [cell.value for cell in wb["订单导入"][1]]
    missing_x = [h for h in REQUIRED_CSV_HEADERS if h not in xheaders]
    if missing_x:
        fail(f"xlsx missing headers: {missing_x}")


def check_raw_push() -> None:
    payload = json.loads(RAW_PUSH.read_text(encoding="utf-8"))
    if not isinstance(payload, list) or len(payload) < 3:
        fail("raw push example must be [shop_no, trades, orders]")
    shop_no, trades, orders = payload[0], payload[1], payload[2]
    if not isinstance(shop_no, str) or not shop_no:
        fail("raw push shop_no must be a non-empty string")
    if not trades or not orders:
        fail("raw push trades/orders cannot be empty")
    trade = trades[0]
    for key in ("tid", "trade_status", "pay_status", "receivable", "receiver_name", "receiver_mobile"):
        if key not in trade:
            fail(f"raw push trade missing {key}")
    if int(trade["order_count"]) != len(orders):
        fail("raw push order_count must equal detail count")
    for order in orders:
        if order.get("tid") != trade["tid"]:
            fail("raw push detail tid mismatch")
        if not order.get("spec_no"):
            fail("raw push detail missing spec_no")
        if float(order["num"]) <= 0:
            fail("raw push detail quantity must be > 0")


def check_trade_import() -> None:
    payload = json.loads(TRADE_IMPORT.read_text(encoding="utf-8"))
    if not isinstance(payload, list) or not payload or not isinstance(payload[0], list):
        fail("trade import example must be [[order, ...]]")
    order = payload[0][0]
    for key in (
        "trade_no",
        "shop_name",
        "warehouse_name",
        "logistics_name",
        "logistics_no",
        "merchant_no",
        "num",
        "receivable",
        "receiver_name",
        "receiver_mobile",
        "receiver_address",
    ):
        if key not in order:
            fail(f"trade import missing {key}")


def check_readme_links() -> None:
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    for rel in (
        "docs/旺店通旗舰版-订单导入操作手册.md",
        "templates/订单导入模板.csv",
        "templates/订单导入模板.xlsx",
        "examples/raw_trade_push_self2.json",
        "examples/trade_import_upload.json",
    ):
        if rel not in readme:
            fail(f"README missing link to {rel}")
        if not (ROOT / rel).is_file():
            fail(f"missing file {rel}")


def main() -> None:
    check_handbook()
    check_csv()
    check_raw_push()
    check_trade_import()
    check_readme_links()
    print("OK: handbook, template, and examples passed validation")


if __name__ == "__main__":
    sys.exit(main())
